import openpyxl
import re
import os
from typing import Tuple, List, Dict, Any

CNPJ_FILIAL = {
    "05.393.606/0001-58": "429",
    "52.803.025/0001-27": "601",
    "71.883.656/0001-48": "Nevine",
}

def parse_vr_sheet(xlsx_path: str, mes_solicitado: str) -> Tuple[str, List[Dict[str, Any]]]:
    if not os.path.exists(xlsx_path):
        raise FileNotFoundError(f"Planilha nao encontrada: {xlsx_path}")

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    mes_clean = mes_solicitado.upper().strip()
    target_sheet = None
    
    # Try exact or partial match with year or month name
    for name in reversed(wb.sheetnames):
        name_upper = name.upper()
        if mes_clean in name_upper or all(p in name_upper for p in mes_clean.split()):
            target_sheet = name
            break

    if not target_sheet:
        # Fallback to last month matching
        for name in reversed(wb.sheetnames):
            if mes_clean.split()[0] in name.upper():
                target_sheet = name
                break

    if not target_sheet:
        raise ValueError(f"Aba correspondente a '{mes_solicitado}' nao encontrada na planilha. Abas: {wb.sheetnames[-10:]}")

    sheet = wb[target_sheet]
    colaboradores = []

    for row in sheet.iter_rows(values_only=True):
        if not row or row[0] is None:
            continue
        col0 = str(row[0]).strip()
        if not col0 or col0.upper() in ["NOME", "TOTAL", "TOTAL GERAL"] or "INFORMA" in col0.upper() or "VALE REFEI" in col0.upper():
            continue

        # Look for CPF in col 1
        col1 = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
        if not col1 or col1.upper() == "CPF":
            continue

        # Look for Valor in col 2
        col2 = row[2] if len(row) > 2 else None
        if not isinstance(col2, (int, float)) or col2 <= 0:
            continue

        # CNPJ in col 4
        col4 = str(row[4]).strip() if len(row) > 4 and row[4] is not None else ""
        cnpj_match = re.search(r"\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}", col4)
        cnpj = cnpj_match.group(0) if cnpj_match else ""
        filial = CNPJ_FILIAL.get(cnpj, "429")

        # Observacoes in col 5
        obs = str(row[5]).strip() if len(row) > 5 and row[5] is not None else ""

        colaboradores.append({
            "nome": col0,
            "cpf": col1,
            "valor": float(col2),
            "cnpj": cnpj,
            "filial": filial,
            "observacoes_planilha": obs
        })

    return target_sheet, colaboradores

if __name__ == "__main__":
    base = os.path.dirname(os.path.dirname(__file__))
    xlsx = os.path.join(base, "Planilha VR.xlsx")
    sheet_name, lista = parse_vr_sheet(xlsx, "OUTUBRO 2026")
    print(f"Sucesso! Aba: {sheet_name}, Total de colaboradores: {len(lista)}")
    total = sum(c['valor'] for c in lista)
    print(f"Valor Total Calculado na Planilha: R$ {total:,.2f}")
    for i, c in enumerate(lista, 1):
        print(f"  {i:02d}. {c['nome']} | Filial: {c['filial']} | R$ {c['valor']:.2f} | Obs: {c['observacoes_planilha']}")
